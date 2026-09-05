import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=50, bottom=50, left=60, right=60):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="444444", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def create_handover_doc(output_path):
    doc = docx.Document()

    # Strict A4 Page Setup with compact margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # Set default style to Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0, 0, 0)

    # 1. Header Table (Quốc hiệu & Tên công ty)
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    cell_left = header_table.cell(0, 0)
    cell_left.width = Inches(3.4)
    p_left = cell_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.paragraph_format.space_before = Pt(0)
    p_left.paragraph_format.space_after = Pt(0)
    p_left.paragraph_format.line_spacing = 1.05
    r1 = p_left.add_run("CÔNG TY TNHH MINH ĐẠO PHÁT\n")
    r1.bold = True
    r1.font.size = Pt(9.5)
    r2 = p_left.add_run("Tổ DP Cam Giá 2, P. Gia Sàng, Tỉnh Thái Nguyên\n")
    r2.font.size = Pt(8)
    r2.font.italic = True
    r3 = p_left.add_run("Số: ........../BBBG-MĐP")
    r3.font.size = Pt(8.5)

    cell_right = header_table.cell(0, 1)
    cell_right.width = Inches(3.5)
    p_right = cell_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.paragraph_format.space_before = Pt(0)
    p_right.paragraph_format.space_after = Pt(0)
    p_right.paragraph_format.line_spacing = 1.05
    r4 = p_right.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r4.bold = True
    r4.font.size = Pt(9.5)
    r5 = p_right.add_run("Độc lập - Tự do - Hạnh phúc\n")
    r5.bold = True
    r5.font.size = Pt(9.5)
    r6 = p_right.add_run("---------------------\n")
    r6.font.size = Pt(7)
    r7 = p_right.add_run("Ngày ..... tháng ..... năm 2026")
    r7.font.size = Pt(8.5)
    r7.font.italic = True

    # 2. Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(1)
    run_title = p_title.add_run("BIÊN BẢN BÀN GIAO PHỤ KIỆN")
    run_title.bold = True
    run_title.font.size = Pt(13)
    run_title.font.color.rgb = RGBColor(0, 32, 96)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(5)
    run_sub = p_sub.add_run("(V/v: Điều chuyển phụ kiện từ Showroom Thuận An sang Showroom Phạm Văn Đồng)")
    run_sub.font.italic = True
    run_sub.font.size = Pt(9.5)

    # 3. Căn cứ & Thông tin chung
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(0)
    p_ref.paragraph_format.space_after = Pt(3)
    p_ref.paragraph_format.line_spacing = 1.05
    r = p_ref.add_run("Căn cứ: ")
    r.bold = True
    r.font.italic = True
    p_ref.add_run("Đơn đề xuất số ")
    r_code = p_ref.add_run("02.202607.01026")
    r_code.bold = True
    p_ref.add_run(" ngày 18/07/2026 về việc xuất điều chuyển lô phụ kiện từ kho Thuận An sang SR Phạm Văn Đồng phục vụ trưng bày.")

    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(3)
    p_info.paragraph_format.line_spacing = 1.05
    p_info.add_run("Hôm nay, ngày ..... tháng ..... năm 2026, tại Showroom, các bên cùng tiến hành bàn giao gồm có:")

    # Side A
    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(2)
    p_a.paragraph_format.line_spacing = 1.05
    r_a = p_a.add_run("I. BÊN GIAO (SR THUẬN AN): ")
    r_a.bold = True
    p_a.add_run("Ông/Bà: .............................................................. Chức vụ: ..............................................................")

    # Side B
    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_before = Pt(0)
    p_b.paragraph_format.space_after = Pt(4)
    p_b.paragraph_format.line_spacing = 1.05
    r_b = p_b.add_run("II. BÊN NHẬN (SR PHẠM VĂN ĐỒNG): ")
    r_b.bold = True
    p_b.add_run("Ông/Bà: .............................................................. Chức vụ: ..............................................................")

    # 4. Table of Items
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="444444", sz="4")

    headers = [
        ("STT", Inches(0.35), WD_ALIGN_PARAGRAPH.CENTER),
        ("Diễn giải / Tên hàng hóa", Inches(2.3), WD_ALIGN_PARAGRAPH.CENTER),
        ("ĐVT", Inches(0.4), WD_ALIGN_PARAGRAPH.CENTER),
        ("SL", Inches(0.4), WD_ALIGN_PARAGRAPH.CENTER),
        ("Giá (VNĐ)", Inches(0.85), WD_ALIGN_PARAGRAPH.CENTER),
        ("Thuế (VNĐ)", Inches(0.75), WD_ALIGN_PARAGRAPH.CENTER),
        ("Thành tiền (VNĐ)", Inches(0.95), WD_ALIGN_PARAGRAPH.CENTER),
        ("Ghi chú", Inches(0.75), WD_ALIGN_PARAGRAPH.CENTER),
    ]

    hdr_cells = table.rows[0].cells
    for i, (text, width, align) in enumerate(headers):
        hdr_cells[i].width = width
        set_cell_background(hdr_cells[i], "EAECEE")
        set_cell_margins(hdr_cells[i], top=50, bottom=50, left=40, right=40)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(8.5)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    items_data = [
        ("1", "Xuất điều chuyển áo polo mã ACS30000001", "cái", "1", "254.584", "20.367", "274.951", "Trưng bày"),
        ("2", "Xuất điều chuyển áo polo mã ACS30000016", "cái", "1", "254.584", "20.367", "274.951", "Trưng bày"),
        ("3", "Xuất điều chuyển ví da mã ACS20000001", "cái", "2", "176.945", "28.311", "382.201", "Trưng bày"),
        ("4", "Xuất điều chuyển ô gập mã ACS20000012", "cái", "3", "188.380", "45.211", "610.351", "Trưng bày"),
        ("5", "Xuất điều chuyển thắt lưng mã ACS30000064", "cái", "1", "2.126.945", "170.156", "2.297.101", "Trưng bày"),
        ("6", "Xuất điều chuyển thảm sàn VF5 mã ACS10000005", "Bộ", "1", "1.185.047", "94.804", "1.279.851", "Trưng bày"),
        ("7", "Xuất điều chuyển tấm che nắng VF6 mã ACS10000037", "cái", "1", "325.000", "26.000", "351.000", "Trưng bày"),
    ]

    for item in items_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(item):
            row_cells[i].width = headers[i][1]
            set_cell_margins(row_cells[i], top=40, bottom=40, left=40, right=40)
            p = row_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i in [0, 2, 3, 7]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i in [4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Total Row
    total_row = table.add_row().cells
    for i in range(8):
        total_row[i].width = headers[i][1]
        set_cell_margins(total_row[i], top=50, bottom=50, left=40, right=40)
        set_cell_background(total_row[i], "F2F4F4")

    total_row[0].merge(total_row[2])
    p_tot_label = total_row[0].paragraphs[0]
    p_tot_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tot_label.paragraph_format.space_before = Pt(0)
    p_tot_label.paragraph_format.space_after = Pt(0)
    r_tot_lbl = p_tot_label.add_run("Tổng cộng:")
    r_tot_lbl.bold = True
    r_tot_lbl.font.size = Pt(8.5)

    p_tot_qty = total_row[3].paragraphs[0]
    p_tot_qty.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tot_qty.paragraph_format.space_before = Pt(0)
    p_tot_qty.paragraph_format.space_after = Pt(0)
    r_tot_qty = p_tot_qty.add_run("10")
    r_tot_qty.bold = True
    r_tot_qty.font.size = Pt(8.5)

    total_row[4].paragraphs[0].text = ""
    total_row[5].paragraphs[0].text = ""

    p_tot_val = total_row[6].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tot_val.paragraph_format.space_before = Pt(0)
    p_tot_val.paragraph_format.space_after = Pt(0)
    r_tot_val = p_tot_val.add_run("5.470.406")
    r_tot_val.bold = True
    r_tot_val.font.size = Pt(8.5)

    total_row[7].paragraphs[0].text = ""

    # 5. Amount in words & Notes
    p_words = doc.add_paragraph()
    p_words.paragraph_format.space_before = Pt(4)
    p_words.paragraph_format.space_after = Pt(3)
    p_words.paragraph_format.line_spacing = 1.05
    r_w1 = p_words.add_run("- Tổng số tiền (bằng chữ): ")
    r_w1.bold = True
    r_w2 = p_words.add_run("Năm triệu bốn trăm bảy mươi nghìn bốn trăm lẻ sáu đồng.")
    r_w2.font.italic = True

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.paragraph_format.line_spacing = 1.05
    p_note.add_run(
        "- Bên Nhận xác nhận đã kiểm tra và nhận đủ số lượng, đúng quy cách, hàng mới 100% để trưng bày tại Showroom Phạm Văn Đồng.\n"
        "- Biên bản này được lập thành 02 bản có giá trị như nhau, mỗi bên giữ 01 bản làm căn cứ theo dõi tài sản."
    )

    # 6. Signatures - 2 columns
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    col_widths = [Inches(3.45), Inches(3.45)]

    sig_headers = [
        ("ĐẠI DIỆN BÊN GIAO\n(Showroom Thuận An)", "(Ký, ghi rõ họ tên)\n\n\n\n\n\n........................................................"),
        ("ĐẠI DIỆN BÊN NHẬN\n(Showroom Phạm Văn Đồng)", "(Ký, ghi rõ họ tên)\n\n\n\n\n\n........................................................"),
    ]

    for col_idx, (title, sub) in enumerate(sig_headers):
        cell = sig_table.cell(0, col_idx)
        cell.width = col_widths[col_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        r_t = p.add_run(title + "\n")
        r_t.bold = True
        r_t.font.size = Pt(10)
        r_s = p.add_run(sub)
        r_s.font.italic = True
        r_s.font.size = Pt(9)

    doc.save(output_path)
    print(f"Document compact A4 successfully created at: {output_path}")

if __name__ == "__main__":
    output_file = r"c:\Users\USER\Documents\ordermanagement\Bien_Ban_Ban_Giao_Phu_Kien_Thuan_An_Pham_Van_Dong.docx"
    create_handover_doc(output_file)
